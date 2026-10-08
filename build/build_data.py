"""
Sell-out 網頁報表：資料建置
讀 sell-out 明細＋NS 維度表（資料包或 NS 匯出 CSV），照 Power Query 01–05 的邏輯拆組，
彙總後以密碼加密（PBKDF2-SHA256 → AES-256-GCM）輸出 docs/data.enc。網頁開啟時輸入同一個密碼解密。

用法：
    python build/build_data.py                 # 密碼讀 build/password.txt（不進版控）或環境變數 SELLOUT_PASSWORD
    python build/build_data.py --check         # 只驗算，不輸出：印出 2026/01 各通路折扣率（網站口徑，兩種 BOM）
路徑預設為上一層 SellOut報表 資料夾，可用 --sellout、--datapack、--nsdir 覆寫。
"""
import argparse
import base64
import csv
import gzip
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

HERE = Path(__file__).resolve().parent
WEB = HERE.parent
ROOT = WEB.parent
PBKDF2_ITER = 250_000


def txt(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def num(v):
    if v is None or v == "":
        return 0.0
    return float(v)


# ---------------------------------------------------------------- 01 F_SO
def read_sellout(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for name in wb.sheetnames:
        if not re.fullmatch(r"\d{4}年度", name):
            continue
        it = wb[name].iter_rows(values_only=True)
        next(it)  # 第 1 列小計
        hdr = [txt(h) for h in next(it)]
        ix = {h: i for i, h in enumerate(hdr) if h}
        for r in it:
            item = txt(r[ix["DRWU品號"]])
            if not item:
                if num(r[ix["台幣金額"]]):
                    print(f"[注意] 工作表「{name}」有品號空白但有金額的列：{r[ix['通路別']]} {r[ix['日期']]} 金額 {num(r[ix['台幣金額']]):,.0f}，已略過")
                continue
            rows.append({
                "ym": txt(r[ix["日期"]]),
                "ch": txt(r[ix["通路別"]]),
                "item": item,
                "qty": num(r[ix["銷售數量"]]),
                "amt": num(r[ix["台幣金額"]]),
                "cur": txt(r[ix["幣別"]]) if "幣別" in ix else None,
            })
    return rows


# ---------------------------------------------------------------- 維度表
def read_table(wb, sheet):
    it = wb[sheet].iter_rows(values_only=True)
    hdr = [txt(h) for h in next(it)]
    return [dict(zip(hdr, r)) for r in it if any(v is not None for v in r)]


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_ch_adjust():
    """通路調整（build/通路類別調整.csv）→ [[通路別, 通路名稱, 通路類別, 說明]]，欄位留空＝不改。
    標題列不對（例：Excel 存檔後欄名被改掉）時直接停止，避免整份調整失效卻沒人發現。"""
    f = HERE / "通路類別調整.csv"
    if not f.exists():
        return []
    with open(f, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    hdr = [h.strip() for h in rows[0]] if rows else []
    if "通路別" not in hdr or not {"通路名稱", "通路類別"} & set(hdr):
        sys.exit(f"[錯誤] {f.name} 的標題列應為「通路別,通路名稱,通路類別,說明」，目前是「{','.join(hdr)}」")
    recs = [dict(zip(hdr, r)) for r in rows[1:]]
    return [[txt(r.get(k)) for k in ("通路別", "通路名稱", "通路類別", "說明")] for r in recs if txt(r.get("通路別"))]


def load_items(pack, nsdir):
    """回傳 {品號: dict}；NS 匯出的 D_Item.csv 存在就用它，否則用資料包的 D_Item。"""
    items = {}
    f = nsdir / "D_Item.csv" if nsdir else None
    if f and f.exists():
        src = []
        for r in read_csv(f):
            src.append({
                "code": txt(r.get("item_code")), "name": txt(r.get("item_name")), "spec": txt(r.get("spec")),
                "p3": txt(r.get("p3")), "p6": txt(r.get("p6")), "series": txt(r.get("series")),
                "dosage": txt(r.get("dosage")), "inv": txt(r.get("inv_type")), "brand": txt(r.get("brand")),
                "price": float(r["list_price"]) if r.get("list_price") not in (None, "") else None,
                "inactive": txt(r.get("inactive")) or "",
            })
        origin = "NS 匯出 D_Item.csv"
    else:
        src = []
        for r in read_table(pack, "D_Item"):
            src.append({
                "code": txt(r["品號"]), "name": txt(r["品名"]), "spec": txt(r["產品代碼"]),
                "p3": txt(r["前三碼"]), "p6": txt(r["前六碼"]), "series": txt(r["系列"]),
                "dosage": txt(r["劑型"]), "inv": txt(r["存貨型態"]), "brand": txt(r["品牌"]),
                "price": float(r["售價(含稅)"]) if r["售價(含稅)"] is not None else None,
                "inactive": txt(r["非作用中"]) or "",
            })
        origin = "資料包 D_Item"
    # 同品號多筆：啟用中、有售價者優先
    src.sort(key=lambda d: (d["code"] or "", str(d["inactive"]).strip().lower() in ("t", "true", "yes", "y", "是", "1"),
                            -(d["price"] if d["price"] is not None else -1)))
    for d in src:
        if d["code"] and d["code"] not in items:
            spec = d["spec"] or ""
            if not d["brand"]:
                d["brand"] = "Redermx" if spec.startswith("X") else "DR.WU"
            core = spec.lstrip("xX").upper()
            if d["inv"] == "組包" and spec and core[:2] in ("PB", "PA", "PT"):
                d["p3"] = d["p6"] = "組合包"
            items[d["code"]] = d
    return items, origin


def extra_bom_pairs():
    f = HERE / "補充BOM.csv"
    if not f.exists():
        return []
    return [(txt(r["組包品號"]), txt(r["成分品號"]), r["用量"] or 0,
             float(r["分攤比例"]) if r.get("分攤比例") not in (None, "") else None, "補充BOM")
            for r in read_csv(f) if txt(r["組包品號"]) and txt(r["成分品號"])]


def load_bom(pack, nsdir, use_temp):
    """回傳 {組包品號: [(成分品號, 用量, 分攤比例, 來源)]}，同 Power Query 02。"""
    def add(dst, rows, skip):
        for kit, comp, q, share, s in rows:
            if kit in skip or share is None or kit is None or comp is None:
                continue
            dst[kit].append((comp, float(q or 0), float(share), s))

    ns = defaultdict(list)
    f = nsdir / "D_BOM.csv" if nsdir else None
    if f and f.exists():
        add(ns, [(txt(r["kit_code"]), txt(r["comp_code"]), r["qty_per_kit"] or 0,
                  float(r["share"]) if r.get("share") not in (None, "") else None,
                  txt(r.get("source")) or "NS BOM") for r in read_csv(f)], set())
    else:
        add(ns, [(txt(r["組包品號"]), txt(r["成分品號"]), r["用量"], r["分攤比例"], txt(r["來源"]) or "NS BOM")
                 for r in read_table(pack, "D_BOM")], set())
    # 補充 BOM（build/補充BOM.csv）：NS 沒有 BOM、業務指定要拆的組包，兩種模式都用
    add(ns, extra_bom_pairs(), set(ns))
    bom = defaultdict(list, {k: list(v) for k, v in ns.items()})
    if use_temp:
        f2 = nsdir / "D_BOM_補.csv" if nsdir else None
        if f2 and f2.exists():
            add(bom, [(txt(r["kit_code"]), txt(r["comp_code"]), r["qty_per_kit"] or 0,
                       float(r["share"]) if r.get("share") not in (None, "") else None,
                       txt(r.get("source")) or "組合單") for r in read_csv(f2)], set(bom))
        add(bom, [(txt(r["組包品號"]), txt(r["成分品號"]), r["用量"], r["分攤比例"], txt(r["來源"]))
                  for r in read_table(pack, "D_BOM_暫用")], set(bom))
    return bom


# ---------------------------------------------------------------- 03 拆組
def explode(fso, bom, items):
    out = []
    for r in fso:
        cur = [(r["item"], r["qty"], r["amt"], False)]
        for _ in range(2):  # 成分本身又是組包時再拆一層
            nxt = []
            for item, q, a, hit in cur:
                comps = bom.get(item)
                if not comps:
                    nxt.append((item, q, a, hit))
                else:
                    nxt.extend((c, q * u, a * s, True) for c, u, s, _ in comps)
            cur = nxt
        for item, q, a, hit in cur:
            it = items.get(item)
            src = "組包未拆" if it and it["inv"] == "組包" else ("拆組" if hit else "單品")
            out.append({"ym": r["ym"], "ch": r["ch"], "orig": r["item"], "item": item, "qty": q, "amt": a, "src": src})
    return out


# ---------------------------------------------------------------- 驗算
def check(fx, items, ch_name):
    agg = defaultdict(lambda: [0.0, 0.0, 0.0])
    for r in fx:
        if r["ym"] != "202601":
            continue
        p = round((items.get(r["item"]) or {}).get("price") or 0)   # 裸瓶用自己的售價（網站同口徑）
        for k in (ch_name.get(r["ch"]) or f"{r['ch']}（未對到通路）", "全通路"):   # 同網頁 prepare()
            a = agg[k]
            a[0] += r["qty"]; a[1] += r["amt"]; a[2] += r["qty"] * p
    return {k: (v[0], v[1], v[2], v[1] / v[2] if v[2] else None) for k, v in agg.items()}


def encrypt(payload: bytes, password: str):
    salt, iv = os.urandom(16), os.urandom(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=PBKDF2_ITER).derive(password.encode())
    ct = AESGCM(key).encrypt(iv, payload, None)
    b64 = lambda b: base64.b64encode(b).decode()
    return {"v": 1, "kdf": "PBKDF2-SHA256", "iter": PBKDF2_ITER, "salt": b64(salt), "iv": b64(iv), "ct": b64(ct)}


def latest_sellout():
    """03_原始資料_業務提供 裡最新的 sell-out 明細檔（檔名含 SellOut、不是彙總表／業績總表；子資料夾的舊版不算）。"""
    folder = ROOT / "03_原始資料_業務提供"
    cands = [f for f in folder.glob("*.xlsx")
             if "sellout" in f.name.lower().replace("-", "") and not any(k in f.name for k in ("彙總", "總表"))
             and not f.name.startswith("~$")]
    if not cands:
        sys.exit(f"[錯誤] {folder} 裡找不到 sell-out 明細檔（檔名要含 SellOut），或用 --sellout 指定")
    return max(cands, key=lambda f: f.stat().st_mtime)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sellout", default=None, help="sell-out 明細檔；預設取 03_原始資料_業務提供 裡最新的一個")
    ap.add_argument("--datapack", default=str(ROOT / "05_資料包" / "SellOut報表資料包_NS對照與拆組_20261007.xlsx"))
    ap.add_argument("--nsdir", default=str(ROOT / "05_資料包" / "NS匯出"))
    ap.add_argument("--out", default=str(WEB / "docs" / "data.enc"))
    ap.add_argument("--check", action="store_true", help="只驗算，不輸出")
    args = ap.parse_args()
    if not args.sellout:
        args.sellout = str(latest_sellout())
    print(f"sell-out 明細：{args.sellout}")

    nsdir = Path(args.nsdir) if Path(args.nsdir).is_dir() else None
    pack = openpyxl.load_workbook(args.datapack, read_only=True, data_only=True)
    fso = read_sellout(args.sellout)
    items, item_origin = load_items(pack, nsdir)
    channels = {txt(r["通路別"]): r for r in read_table(pack, "D_Channel")}
    ch_name = {k: txt(v["通路名稱"]) for k, v in channels.items()}
    # 通路調整（build/通路類別調整.csv）：原樣放進 data.enc（chAdjust），由網頁端 buildModel 套用，
    # 網頁上傳新的通路對照時調整也會保留；channels 維持資料包原值（備註、地區不被覆蓋）。這裡只套用通路名稱給下方驗算
    ch_adjust = read_ch_adjust()
    for ch, name, _, _ in ch_adjust:
        if ch not in channels:
            print(f"[注意] 通路類別調整.csv 的通路別「{ch}」不在 D_Channel，沒有套用（請檢查是否打錯）")
        elif name:
            ch_name[ch] = name

    fx_ns = explode(fso, load_bom(pack, nsdir, False), items)
    fx_tmp = explode(fso, load_bom(pack, nsdir, True), items)

    # ---- 驗算（UseTempBOM＝false 對開發文檔第 7 節「重算」欄）
    print(f"F_SO {len(fso)} 列，金額 {sum(r['amt'] for r in fso):,.0f}；D_Item 來源：{item_origin}")
    print(f"拆組前後金額差：NS {sum(r['amt'] for r in fx_ns) - sum(r['amt'] for r in fso):.6f}"
          f"／含暫用 {sum(r['amt'] for r in fx_tmp) - sum(r['amt'] for r in fso):.6f}")
    c_ns, c_tmp = check(fx_ns, items, ch_name), check(fx_tmp, items, ch_name)
    print("2026/01 本期折扣率（網站口徑：裸瓶用裸瓶售價）  只用NS BOM（預設）   含暫用BOM")
    for k in sorted(c_ns, key=lambda k: -(c_ns[k][3] or 0)):
        print(f"  {k:<10} {c_ns[k][3] or 0:>16.2%} {c_tmp.get(k, (0, 0, 0, 0))[3] or 0:>14.2%}")
    if args.check:
        return

    # ---- 輸出（v2）：維度表原樣＋未拆組明細；拆組由網頁端執行（與上傳更新共用同一套程式）
    def bom_rows(pairs):
        out = []
        for kit, comp, q, share in pairs:
            if kit and comp and share is not None:
                out.append([kit, comp, float(q or 0), float(share)])
        return out

    f_ns, f_bu = (nsdir / "D_BOM.csv", nsdir / "D_BOM_補.csv") if nsdir else (None, None)
    if f_ns and f_ns.exists():
        ns_pairs = [(txt(r["kit_code"]), txt(r["comp_code"]), r["qty_per_kit"],
                     float(r["share"]) if r.get("share") not in (None, "") else None) for r in read_csv(f_ns)]
    else:
        ns_pairs = [(txt(r["組包品號"]), txt(r["成分品號"]), r["用量"], r["分攤比例"]) for r in read_table(pack, "D_BOM")]
    bu_pairs = [(txt(r["kit_code"]), txt(r["comp_code"]), r["qty_per_kit"],
                 float(r["share"]) if r.get("share") not in (None, "") else None)
                for r in read_csv(f_bu)] if f_bu and f_bu.exists() else []
    fb_pairs = [(txt(r["組包品號"]), txt(r["成分品號"]), r["用量"], r["分攤比例"]) for r in read_table(pack, "D_BOM_暫用")]

    so = defaultdict(lambda: [0.0, 0.0])
    for r in fso:
        a = so[(r["ym"], r["ch"], r["item"], r["cur"] or "")]
        a[0] += r["qty"]; a[1] += r["amt"]
    data = {
        "v": 2,
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source": {"sellout": Path(args.sellout).name, "datapack": Path(args.datapack).name, "item": item_origin},
        "itemCols": ["品號", "品名", "產品代碼", "前三碼", "前六碼", "系列", "劑型", "存貨型態", "品牌", "售價", "非作用中"],
        "items": [[c, d["name"], d["spec"], d["p3"], d["p6"], d["series"], d["dosage"], d["inv"], d["brand"],
                   round(d["price"]) if d["price"] is not None else None,
                   "T" if str(d["inactive"]).strip().lower() in ("t", "true", "yes", "y", "是", "1") else "F"]
                  for c, d in sorted(items.items())],
        "bom": {"ns": bom_rows(ns_pairs), "bu": bom_rows(bu_pairs), "fb": bom_rows(fb_pairs),
                "extra": bom_rows([p[:4] for p in extra_bom_pairs()])},
        "channels": [[txt(r["通路別"]), txt(r["通路名稱"]), txt(r["通路類別"]), txt(r["地區"]), txt(r["備註"])]
                     for r in channels.values()],
        "chAdjust": ch_adjust,
        "so": [[ym, ch, it, cur, round(v[0], 4), round(v[1], 4)] for (ym, ch, it, cur), v in sorted(so.items())],
        "seriesOk": [r["系列確認品號"].strip() for r in read_csv(HERE / "系列確認.csv") if r["系列確認品號"].strip()]
                    if (HERE / "系列確認.csv").exists() else [],
        "bareMap": [[r["裸瓶品號"].strip(), r["正貨品號"].strip()] for r in read_csv(HERE / "裸瓶對照.csv")
                    if r["裸瓶品號"].strip() and r["正貨品號"].strip()] if (HERE / "裸瓶對照.csv").exists() else [],
        "seriesMap": [[txt(r["原系列名稱"]), txt(r["報表系列名稱"])] for r in read_csv(HERE / "系列對照.csv")
                      if txt(r["原系列名稱"]) and txt(r["報表系列名稱"])] if (HERE / "系列對照.csv").exists() else [],
    }
    raw = json.dumps(data, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    pw = os.environ.get("SELLOUT_PASSWORD")
    pwf = HERE / "password.txt"
    if not pw and pwf.exists():
        pw = pwf.read_text(encoding="utf-8-sig").strip()   # 記事本存成含 BOM 也能讀
    if not pw:
        sys.exit("沒有密碼：請建立 build/password.txt 或設定環境變數 SELLOUT_PASSWORD")
    enc = encrypt(gzip.compress(raw, 9), pw)
    Path(args.out).write_text(json.dumps(enc), encoding="utf-8")
    print(f"已輸出 {args.out}（JSON {len(raw)/1e6:.2f} MB → 加密 {Path(args.out).stat().st_size/1e6:.2f} MB）")


if __name__ == "__main__":
    main()
